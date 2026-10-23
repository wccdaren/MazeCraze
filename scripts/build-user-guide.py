"""Build the illustrated Word manual from its editable Markdown source.

Requires python-docx==1.2.0. Page-number mapping is optional; see docs/guide-pages.json.
This script changes documentation only, never game code or configuration.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/Map-Editor-User-Guide.md'
TARGET = ROOT / 'docs/Maze-Craze-Map-Editor-User-Guide.docx'
GREEN = '214C3D'
PALE = 'EEF4E9'
INK = '26332E'
MUTED = '57665D'


def shade(element, fill):
    props = element.get_or_add_tcPr() if element.tag == qn('w:tc') else element.get_or_add_pPr()
    existing = props.find(qn('w:shd'))
    if existing is not None:
        props.remove(existing)
    item = OxmlElement('w:shd')
    item.set(qn('w:fill'), fill)
    props.append(item)


def rich(paragraph, text):
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor.from_string(GREEN)
        else:
            paragraph.add_run(part)


def simple_field(paragraph, instruction, result, size=11):
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), instruction)
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Calibri')
    fonts.set(qn('w:hAnsi'), 'Calibri')
    points = OxmlElement('w:sz')
    points.set(qn('w:val'), str(round(size * 2)))
    properties.extend([fonts, points])
    text = OxmlElement('w:t')
    text.text = str(result)
    run.extend([properties, text])
    field.append(run)
    paragraph._p.append(field)


def anchor(paragraph, name, bookmark_id):
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bookmark_id))
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(bookmark_id))
    # Paragraph properties must remain the first child in valid WordprocessingML.
    paragraph._p.insert(1 if paragraph._p.find(qn('w:pPr')) is not None else 0, start)
    paragraph._p.append(end)


def contents_link(paragraph, title, name):
    link = OxmlElement('w:hyperlink')
    link.set(qn('w:anchor'), name)
    link.set(qn('w:history'), '1')
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), GREEN)
    props.append(color)
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Calibri')
    fonts.set(qn('w:hAnsi'), 'Calibri')
    props.insert(0, fonts)
    run.append(props)
    text = OxmlElement('w:t')
    text.text = title
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def setup_styles(document):
    styles = document.styles
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for level, size in [(1, 25), (2, 16), (3, 12.5)]:
        style = styles[f'Heading {level}']
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(GREEN)
        style.paragraph_format.space_before = Pt(14 if level > 1 else 0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = level == 1
    for name, size in [('Title', 34), ('Subtitle', 16)]:
        styles[name].font.name = 'Calibri'
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(GREEN)
    for name in ['List Bullet', 'List Number']:
        styles[name].font.name = 'Calibri'
        styles[name].font.size = Pt(11)
        styles[name].paragraph_format.space_after = Pt(5)
    for name in ['Guide Code', 'Guide Callout', 'Guide Caption', 'Guide Contents', 'Guide Small']:
        if name not in styles:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        styles[name].base_style = styles['Normal']
        styles[name].font.name = 'Calibri'
    code = styles['Guide Code']
    code.font.name = 'Consolas'
    code.font.size = Pt(9)
    code.paragraph_format.left_indent = Inches(.12)
    code.paragraph_format.right_indent = Inches(.12)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.05
    callout = styles['Guide Callout']
    callout.font.size = Pt(10.5)
    callout.paragraph_format.left_indent = Inches(.12)
    callout.paragraph_format.right_indent = Inches(.12)
    callout.paragraph_format.space_before = Pt(7)
    callout.paragraph_format.space_after = Pt(10)
    caption = styles['Guide Caption']
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_after = Pt(11)
    caption.paragraph_format.keep_with_next = False
    small = styles['Guide Small']
    small.font.size = Pt(9.5)
    small.font.color.rgb = RGBColor.from_string(MUTED)
    contents = styles['Guide Contents']
    contents.font.size = Pt(11)
    contents.paragraph_format.space_after = Pt(8)
    contents.paragraph_format.tab_stops.add_tab_stop(Inches(6.85), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def table(document, rows):
    if not rows:
        return
    count = len(rows[0])
    if count != 3:
        raise ValueError(f'Expected an examples column in each 3-column table; got {count}.')
    result = document.add_table(rows=0, cols=count)
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    widths = [1.5, 3.0, 2.4]
    for col, width in zip(result.columns, widths):
        col.width = Inches(width)
    for index, source in enumerate(rows):
        cells = result.add_row().cells
        for cell, width, value in zip(cells, widths, source):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.line_spacing = 1.08
            rich(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                if index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            shade(cell._tc, GREEN if index == 0 else (PALE if index % 2 else 'F8FAF5'))
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                part = OxmlElement(f'w:{side}')
                part.set(qn('w:w'), '90')
                part.set(qn('w:type'), 'dxa')
                margins.append(part)
            tc_pr.append(margins)
        row_props = result.rows[-1]._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        row_props.append(no_split)
        if index == 0:
            repeat = OxmlElement('w:tblHeader')
            row_props.append(repeat)
            for cell in cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def picture(document, path, caption):
    from docx.image.image import Image
    image = Image.from_file(str(path))
    max_width, max_height = 6.65, 6.0
    ratio = image.px_height / image.px_width
    width = min(max_width, max_height / ratio)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set('descr', caption)
    document.add_paragraph(caption, 'Guide Caption')


def build(pages):
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.7)
    section.left_margin = Inches(.8)
    section.right_margin = Inches(.8)
    section.header_distance = Inches(.28)
    section.footer_distance = Inches(.28)
    section.different_first_page_header_footer = True
    setup_styles(document)
    header = section.header.paragraphs[0]
    header.add_run('MAZE CRAZE  /  MAP STUDIO').bold = True
    header.add_run('                                      THE MAPMAKER’S HANDBOOK')
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GREEN)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Maze Craze Map Studio  •  User Guide  •  ')
    simple_field(footer, 'PAGE', '1', 8.5)
    for run in footer.runs:
        run.font.size = Pt(8.5)
        run.font.name = 'Calibri'
        run.font.color.rgb = RGBColor.from_string(MUTED)
    document.core_properties.title = 'Maze Craze Map Studio — The Mapmaker’s Handbook'
    document.core_properties.subject = 'Illustrated user manual for seventh-grade mapmakers'
    document.core_properties.author = 'Maze Craze'
    document.core_properties.keywords = 'Maze Craze, map editor, Sebastian, coordinates, puzzles, user guide'
    document.core_properties.comments = 'Documents integrated playtesting and explicit PASS/FAIL item costs. Documentation only.'
    update_fields = OxmlElement('w:updateFields')
    update_fields.set(qn('w:val'), 'true')
    document.settings.element.append(update_fields)

    document.add_paragraph('A SEBASTIAN ADVENTURE', 'Subtitle')
    document.add_paragraph('Maze Craze\nMap Studio', 'Title')
    document.add_paragraph('The Mapmaker’s Handbook', 'Subtitle')
    document.add_paragraph('Build worlds. Write clues. Make every square count.')
    cover = document.add_paragraph()
    cover.paragraph_format.space_before = Pt(10)
    cover.add_run().add_picture(str(ROOT / 'assets/gnome.png'), height=Inches(1.65))
    document.add_paragraph('A complete, illustrated guide for Sebastian and fellow seventh-grade mapmakers.')
    for text in ['A guided first map you can build and play', 'Every editor field, valid settings, examples, and creative uses', 'Recipes for keys, doors, chests, equipment, secrets, and travel', 'The file relationships that make the whole adventure work']:
        document.add_paragraph(text, 'List Bullet')
    document.add_paragraph('Playtesting edition • October 2026', 'Guide Small')
    document.add_paragraph('Offline editing and export • Integrated playtests use the same HTTP(S) game website.', 'Guide Small')
    document.add_page_break()
    title = document.add_paragraph('Contents', 'Title')
    title.paragraph_format.space_after = Pt(12)
    document.add_paragraph('Start with Chapters 1–4. Then use the field reference and recipes as you create.', 'Guide Small')
    source = SOURCE.read_text(encoding='utf-8')
    chapters = re.findall(r'^# (.+)$', source, re.M)
    for index, chapter in enumerate(chapters, 1):
        paragraph = document.add_paragraph(style='Guide Contents')
        contents_link(paragraph, chapter, f'chapter_{index}')
        paragraph.add_run('\t')
        simple_field(paragraph, f'PAGEREF chapter_{index} \\h', pages.get(str(index), '…'))
    document.add_paragraph('In Word, use View → Navigation Pane to jump to headings. The contents titles are clickable.', 'Guide Small')
    document.add_paragraph('Field entries use the same pattern: what it does, valid settings, examples, creative ideas, and connections to other settings.', 'Guide Small')

    lines = source.splitlines()
    index = 0
    chapter_index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        heading = re.match(r'^(#{1,3}) (.+)$', line)
        if heading:
            level = len(heading[1])
            paragraph = document.add_heading(heading[2], level)
            if level == 1:
                chapter_index += 1
                anchor(paragraph, f'chapter_{chapter_index}', chapter_index)
            index += 1
            continue
        image = re.match(r'^!\[(.*?)\]\((.*?)\)$', line)
        if image:
            picture(document, SOURCE.parent / image[2], image[1])
            index += 1
            continue
        if line.startswith('```'):
            block = []
            index += 1
            while index < len(lines) and not lines[index].startswith('```'):
                block.append(lines[index])
                index += 1
            for position, code in enumerate(block):
                paragraph = document.add_paragraph(code, 'Guide Code')
                shade(paragraph._p, 'F0F3ED')
                paragraph.paragraph_format.keep_with_next = len(block) < 13 and position < len(block) - 1
            document.add_paragraph().paragraph_format.space_after = Pt(1)
            index += 1
            continue
        if line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].startswith('|'):
                row = [part.strip() for part in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', part) for part in row):
                    rows.append(row)
                index += 1
            table(document, rows)
            continue
        if line.startswith('- '):
            rich(document.add_paragraph(style='List Bullet'), line[2:])
            index += 1
            continue
        if line.startswith('> '):
            paragraph = document.add_paragraph(style='Guide Callout')
            shade(paragraph._p, PALE)
            rich(paragraph, line[2:])
            index += 1
            continue
        paragraph_lines = [line]
        index += 1
        while index < len(lines) and lines[index].strip() and not re.match(r'^(#{1,3} |\||- |> |```|!\[)', lines[index]):
            paragraph_lines.append(lines[index])
            index += 1
        rich(document.add_paragraph(), ' '.join(paragraph_lines))

    document.save(TARGET)
    print(f'Built {TARGET} ({TARGET.stat().st_size:,} bytes)')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--pages', type=Path, help='JSON map of chapter number to rendered page number')
    args = parser.parse_args()
    build(json.loads(args.pages.read_text()) if args.pages else {})
